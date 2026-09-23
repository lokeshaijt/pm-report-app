"""
Pending Purchase Order parser.

Two formats have been seen from the source system:
  1. Curated "Outstanding Qty" sheet (most common): Document No, Order Date,
     Customer Code, Customer Name, Item No, Item Description, Delivery Date,
     Quantity, Outstanding Quantity, Unit Cost - 10 columns, no Location Code.
     This appears to already be pre-scoped to PM-relevant purchases; every row
     is used as-is.
  2. Full company-wide raw extract (seen once, 18-Sep-2026): 18 columns
     including Location Code, mixed with completely unrelated company
     purchases (steel pipes, printer parts, etc). When this format appears,
     only rows with Location Code == "CBEPM" are used - every other location
     is a different plant/department and must be excluded.

The parser detects which format it has by header content, not by sheet name
or position, and applies the right column mapping and filter automatically.
"""
import openpyxl
from collections import defaultdict
from .utils import norm, to_float
from .weeks import parse_date

CBEPM_LOCATION = "CBEPM"


def _find_po_sheet_and_format(wb):
    """
    Returns (ws, format, header_row, col_map) where format is 'curated' or 'raw'.
    col_map keys: item_desc, delivery_date, outstanding, location_code (raw only).
    """
    for name in wb.sheetnames:
        ws = wb[name]
        # curated format: header on row 1
        header1 = [ws.cell(1, c).value for c in range(1, 11)]
        if header1[5] == "Item Description" and header1[6] == "Delivery Date":
            return ws, "curated", 1, {"item_desc": 6, "delivery_date": 7, "outstanding": 9}
        # raw format: header row varies (commonly row 13), scan for it
        for r in range(1, 20):
            header = [ws.cell(r, c).value for c in range(1, 19)]
            if "Item Description" in header and "Delivery Date" in header and "Location Code" in header:
                col_map = {
                    "item_desc": header.index("Item Description") + 1,
                    "delivery_date": header.index("Delivery Date") + 1,
                    "outstanding": header.index("Outstanding Quantity") + 1,
                    "location_code": header.index("Location Code") + 1,
                }
                return ws, "raw", r, col_map
    # fall back to the last sheet, assume curated layout at row 1
    ws = wb[wb.sheetnames[-1]]
    return ws, "curated", 1, {"item_desc": 6, "delivery_date": 7, "outstanding": 9}


def parse_pending_po(file_obj):
    """
    Returns:
      total_pending_po: {item_key: total_outstanding_qty}
      rows: list of (item_key, delivery_date, qty)   for weekly/monthly bucketing
      format_used: 'curated' or 'raw' (for surfacing to the user)
      n_excluded_other_location: count of rows skipped due to non-CBEPM location (raw format only)
    """
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws, fmt, header_row, col_map = _find_po_sheet_and_format(wb)

    total_pending_po = defaultdict(float)
    rows = []
    n_excluded = 0

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row is None:
            continue
        item_desc = row[col_map["item_desc"] - 1]
        if item_desc is None:
            continue
        if fmt == "raw":
            loc = row[col_map["location_code"] - 1]
            if loc != CBEPM_LOCATION:
                n_excluded += 1
                continue
        outstanding = row[col_map["outstanding"] - 1]
        if outstanding is None:
            continue
        qty = to_float(outstanding)
        if qty == 0:
            continue
        key = norm(item_desc)
        total_pending_po[key] += qty
        d = parse_date(row[col_map["delivery_date"] - 1])
        rows.append((key, d, qty))

    return dict(total_pending_po), rows, fmt, n_excluded

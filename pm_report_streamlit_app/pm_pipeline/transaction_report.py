"""
Optional Transaction Report parser.

Accepts the raw Transaction Report export, either as the native .xlsb
("Transaction Report-DD-Mon-YY_HH-MM-SS.xlsb") or as an .xlsx with the same
24-column layout. The real header row (after a handful of report-filter
summary rows) is located by content, not by a fixed row number.

Builds per (Doc_Date, Item_Name) rows with Floor Wastage / Store Wastage /
Writeoff, using only the STOCKOUT leg of each transaction (every WASTAGE and
WRITEOFF transaction has a paired STOCKIN leg for the same qty - counting
both would double the total):

  Transaction Type == "WASTAGE"  -> wastage, split by From Warehouse:
      From Warehouse in {WH-PM-MDKCBE, WH-PM-MDKCBE2} -> Store Wastage
      every other From Warehouse                      -> Floor Wastage
  Transaction Type == "WRITEOFF" -> Writeoff (any warehouse)

Total Wastage is always Floor Wastage + Store Wastage.
"""
from datetime import date, datetime, timedelta

import openpyxl

from .utils import norm_disp
from .weeks import week_number_for

STORE_FROM_WAREHOUSES = {"WH-PM-MDKCBE", "WH-PM-MDKCBE2"}
REQUIRED_HEADERS = ("Transaction Type", "Item_Name", "Movement Type", "Doc_Date", "Qty")
XLSB_EPOCH = date(1899, 12, 30)


def _xlsb_date(v):
    if not isinstance(v, (int, float)):
        return None
    return XLSB_EPOCH + timedelta(days=int(v))


def _rows_from_xlsx(file_obj):
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = None
    for row in rows_iter:
        cells = {str(c).strip() for c in row if c is not None}
        if all(h in cells for h in REQUIRED_HEADERS):
            header = row
            break
    if header is None:
        raise ValueError("Transaction Report: could not find the expected header row.")
    col = {str(name).strip(): i for i, name in enumerate(header) if name}
    for row in rows_iter:
        if row is None or all(v is None for v in row):
            continue
        yield {k: row[i] for k, i in col.items()}


def _rows_from_xlsb(file_obj):
    import pyxlsb

    with pyxlsb.open_workbook(file_obj) as wb:
        with wb.get_sheet(wb.sheets[0]) as sheet:
            rows_iter = sheet.rows()
            header = None
            for sparse_row in rows_iter:
                cells = {str(c.v).strip() for c in sparse_row if c.v is not None}
                if all(h in cells for h in REQUIRED_HEADERS):
                    header = [c.v for c in sparse_row]
                    break
            if header is None:
                raise ValueError("Transaction Report: could not find the expected header row.")
            col = {str(name).strip(): i for i, name in enumerate(header) if name}
            for sparse_row in rows_iter:
                values = [c.v for c in sparse_row]
                if not values or all(v is None for v in values):
                    continue
                row = {k: (values[i] if i < len(values) else None) for k, i in col.items()}
                row["Doc_Date"] = _xlsb_date(row.get("Doc_Date"))
                yield row


def parse_transaction_report(file_obj):
    """
    Returns a list of dicts sorted by (date, item name), one per
    (Doc_Date, Item_Name) combination that has nonzero activity:
      {date, item_name, uom, item_type,
       floor_wastage, store_wastage, writeoff}
    """
    name = (getattr(file_obj, "name", "") or "").lower()
    file_obj.seek(0)
    rows = _rows_from_xlsb(file_obj) if name.endswith(".xlsb") else _rows_from_xlsx(file_obj)

    agg = {}
    item_uom, item_type = {}, {}

    for r in rows:
        txn_type = str(r.get("Transaction Type") or "").strip().upper()
        if txn_type not in ("WASTAGE", "WRITEOFF"):
            continue
        if str(r.get("Movement Type") or "").strip().upper() != "STOCKOUT":
            continue

        item_name = r.get("Item_Name")
        if not item_name:
            continue
        item_name = norm_disp(item_name)

        d = r.get("Doc_Date")
        if isinstance(d, datetime):
            d = d.date()
        if not isinstance(d, date):
            continue

        try:
            qty = float(r.get("Qty"))
        except (TypeError, ValueError):
            continue

        item_uom.setdefault(item_name, r.get("UOM_Name") or "")
        item_type.setdefault(item_name, r.get("Item_Type_Name") or "")

        key = (d, item_name)
        if key not in agg:
            agg[key] = {"floor": 0.0, "store": 0.0, "writeoff": 0.0}

        if txn_type == "WRITEOFF":
            agg[key]["writeoff"] += qty
        elif str(r.get("From Warehouse") or "").strip() in STORE_FROM_WAREHOUSES:
            agg[key]["store"] += qty
        else:
            agg[key]["floor"] += qty

    out = []
    for (d, item_name), v in sorted(agg.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        out.append({
            "date": d, "item_name": item_name,
            "uom": item_uom.get(item_name, ""), "item_type": item_type.get(item_name, ""),
            "floor_wastage": round(v["floor"], 4), "store_wastage": round(v["store"], 4),
            "writeoff": round(v["writeoff"], 4),
        })
    return out


def filter_last_n_completed_weeks(rows, today, n=6):
    """
    Keeps only rows whose date falls in the N most recently *completed*
    weeks before today's own (in-progress) week - e.g. on a day in week 40,
    "last 6 weeks" means weeks 34-39, not weeks 35-40.
    """
    current_week = week_number_for(today)
    min_week, max_week = current_week - n, current_week - 1
    return [r for r in rows if min_week <= week_number_for(r["date"]) <= max_week]

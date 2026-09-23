"""
Stock Report parser.

The source file's layout has varied run to run:
  - Sometimes a raw dump sheet (Location_Name, Warehouse_Name, Item_Name, ..., Qty Total)
  - Sometimes a ready-made pivot, headed either "Row Labels" or "Item_Name" in
    column A, then one column per warehouse, then Grand Total
    - sometimes with all 4 target warehouses as columns, sometimes with only
      the ones that have nonzero stock (missing columns just mean 0 for that WH)

Current stock is always: sum of on-hand quantity across exactly these 4 locations,
which for a pivot sheet is simply its own Grand Total column.
"""
import openpyxl
from collections import defaultdict
from .utils import norm

TARGET_LOCATIONS = {"WH-PM-MDKCBE", "WH-PM-MDKCBE2", "WH-PZ1-MDKCBE", "WH-PZ1-MDKCBE2"}


def _find_header_row(ws, max_scan=15):
    for r in range(1, max_scan + 1):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        if vals[0] == "Location_Name":
            return r, "raw"
        if vals[0] in ("Row Labels", "Item_Name"):
            return r, "pivot"
    return None, None


def parse_stock(file_obj) -> dict:
    """Returns {normalized_item_name: current_stock_float}.

    Only ONE representation of stock is used, even if the workbook contains
    several sheets that both describe the same underlying stock (e.g. a raw
    dump sheet AND a pivot sheet). Using more than one and summing them
    double-counts every item.
    Priority: any raw dump sheet found (most reliable, since it's filtered
    to the 4 target warehouses explicitly) beats any pivot sheet.
    """
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)

    raw_sheet_info = None   # (sheet_name, header_row)
    pivot_sheet_info = None  # (sheet_name, header_row) - first one found only

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        header_row, kind = _find_header_row(ws)
        if header_row is None:
            continue
        if kind == "raw" and raw_sheet_info is None:
            raw_sheet_info = (sheet_name, header_row)
        elif kind == "pivot" and pivot_sheet_info is None:
            pivot_sheet_info = (sheet_name, header_row)

    stock = defaultdict(float)

    if raw_sheet_info is not None:
        sheet_name, header_row = raw_sheet_info
        ws = wb[sheet_name]
        header = [ws.cell(header_row, c).value for c in range(1, 8)]
        wh_col = header.index("Warehouse_Name") + 1
        name_col = header.index("Item_Name") + 1
        qty_col = header.index("Qty Total") + 1
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            wh = row[wh_col - 1]
            name = row[name_col - 1]
            qty = row[qty_col - 1]
            if wh in TARGET_LOCATIONS and name:
                stock[norm(name)] += (qty or 0)
        return dict(stock)

    if pivot_sheet_info is not None:
        sheet_name, header_row = pivot_sheet_info
        ws = wb[sheet_name]
        headers = []
        c = 1
        while True:
            v = ws.cell(header_row, c).value
            if v is None and c > 1:
                break
            headers.append(v)
            c += 1
        try:
            grand_total_col = headers.index("Grand Total") + 1
        except ValueError:
            grand_total_col = len(headers)
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            name = row[0]
            if not name or str(name).strip().lower() == "grand total":
                continue
            stock[norm(name)] = row[grand_total_col - 1] or 0
        return dict(stock)

    return dict(stock)  # nothing found - empty, will surface as all-zero stock
